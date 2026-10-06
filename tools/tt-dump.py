# tt-dump.py — выгрузка содержимого .xlsx (ТТ) в текстовые дампы для чтения.
# Запуск: python analiz-tep/tools/tt-dump.py "<файл.xlsx>" "<папка вывода>"
import sys, os, json
from openpyxl import load_workbook

src = sys.argv[1]
outdir = sys.argv[2]
os.makedirs(outdir, exist_ok=True)

wb = load_workbook(src, data_only=True)
summary = []
for ws in wb.worksheets:
    rows = list(ws.iter_rows())
    # карта объединённых ячеек: значение верхней левой ячейки на весь диапазон
    merged = {}
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = v
    lines = []
    nonempty = 0
    for r in range(1, ws.max_row + 1):
        cells = []
        for c in range(1, ws.max_column + 1):
            v = merged.get((r, c), ws.cell(r, c).value)
            if v is None:
                continue
            s = str(v).replace('\r\n', '\n').strip()
            if s == '':
                continue
            cells.append('[%s] %s' % (ws.cell(r, c).coordinate if (r, c) not in merged else ws.cell(r, c).coordinate, s))
            nonempty += 1
        if cells:
            lines.append('R%d: %s' % (r, ' | '.join(cells)))
    name = ws.title
    safe = ''.join(ch if ch.isalnum() or ch in ' -_.' else '_' for ch in name)
    with open(os.path.join(outdir, safe + '.txt'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    summary.append({'sheet': name, 'dims': ws.dimensions, 'rows': ws.max_row, 'cols': ws.max_column,
                    'nonempty_cells': nonempty, 'merged': len(ws.merged_cells.ranges),
                    'dump_lines': len(lines), 'file': safe + '.txt'})

with open(os.path.join(outdir, '_summary.json'), 'w', encoding='utf-8') as f:
    json.dump(summary, f, ensure_ascii=False, indent=2)

print('файл: %s' % src)
for s in summary:
    print('  лист %-28s %-12s строк=%-5d столбцов=%-3d непустых=%-5d объединений=%-4d дамп=%s' %
          (s['sheet'], s['dims'], s['rows'], s['cols'], s['nonempty_cells'], s['merged'], s['dump_lines']))
