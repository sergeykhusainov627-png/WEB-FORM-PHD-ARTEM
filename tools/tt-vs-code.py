# tt-vs-code.py — сверка ТТ (xlsx) с кодом веб-формы (backup.txt).
# Запуск: python analiz-tep/tools/tt-vs-code.py "<ТТ.xlsx>" "<модуль.fore>" "<папка вывода>"
import sys, os, re
from openpyxl import load_workbook

tt_path, fore_path, outdir = sys.argv[1], sys.argv[2], sys.argv[3]
os.makedirs(outdir, exist_ok=True)
wb = load_workbook(tt_path, data_only=True)

def col(ws, letter, r1, r2):
    return [ws['%s%d' % (letter, r)] .value for r in range(r1, r2 + 1)]

def norm(v):
    return ('' if v is None else str(v)).strip()

out = []
def p(s=''):
    out.append(s)

# ---------- 1. ТТ ----------
ws_f = wb['Разделы СЭ - Условия выборки']
filters = []          # (раздел, поле, справочник)
section = ''
for r in range(4, ws_f.max_row + 1):
    a, b, c, d = (norm(ws_f.cell(r, i).value) for i in (1, 2, 3, 4))
    if a == 'Раздел':
        section = b
    elif c:
        filters.append((section, c, d))

ws_o = wb['СЭ - поля для вывода']
outputs, osect = [], ''
for r in range(4, ws_o.max_row + 1):
    a, b, c = (norm(ws_o.cell(r, i).value) for i in (1, 2, 3))
    if a and not b:
        osect = a.rstrip(':')
    elif b:
        outputs.append((osect, b, c))

ws_r = wb['Вывод полей в отчете']
report_fields = []    # (поле отчёта, условие, источник, поле-источник, код?, источник кода, поле кода)
for r in range(5, ws_r.max_row + 1):
    v = [norm(ws_r.cell(r, i).value) for i in range(1, 8)]
    if v[0]:
        report_fields.append(tuple(v))

ws_s = wb['Селекционный экран']
screen = []
for r in range(6, ws_s.max_row + 1):
    a, b, c = (norm(ws_s.cell(r, i).value) for i in (16, 17, 18))  # P,Q,R
    if a:
        screen.append((a, b, c))

# ---------- 2. Код ----------
src = open(fore_path, encoding='utf-8').read()
cb_fields = re.findall(r'fillFieldListByStatusFlag\(\s*([A-Za-z_]\w*)\s*,\s*"([^"]+)"\s*\)', src)
params = re.findall(r'UpdateOpenDefHlinkFromSelection\(\s*"([^"]+)"\s*,\s*([A-Za-z_]\w*)\.Selection\s*,\s*"([^"]+)"', src)
direct = re.findall(r'_AnalyzTepHyperlink\.SetParamValue\(\s*C_PARAM_(\w+)\s*,', src)
declared = re.findall(r'^\t(\w+):\s*(Web\w+|IStringList|ITabHyperlink|IMETABASE|Variant)\s*;', src, re.M)

param_names = [q[0] for q in params]
field_names = [f[1] for f in cb_fields]
tt_ids = set()
for rf in report_fields:
    for x in (rf[3], rf[6]):
        for tok in re.findall(r'[A-Z][A-Z0-9_]{2,}', x.upper()):
            tt_ids.add(tok)
for _, _, d in filters:
    if d and d.upper() == d and re.fullmatch(r'[A-Z0-9_]+', d):
        tt_ids.add(d)

# ---------- 3. Отчёт ----------
p('# Сверка ТТ ↔ код веб-формы (машинная выгрузка)')
p()
p('## 0. Сводка')
p('- Полей «Условия выборки» в ТТ: %d; полей «Поля для вывода»: %d; строк «Вывод полей в отчёте»: %d' %
  (len(filters), len(outputs), len(report_fields)))
p('- В коде: параметров гиперссылки (комбо): %d; записей в P_FIELD_LIST (флажки): %d' % (len(param_names), len(field_names)))
p()

p('## 1. «Условия выборки» (ТТ) ↔ комбобоксы формы')
p()
p('| Раздел | Поле на экране | Справочник (ТТ) | Комбо в форме | Атрибут | Статус |')
p('|---|---|---|---|---|---|')
by_param = {q[0]: (q[1], q[2]) for q in params}
for sect, name, dic in filters:
    hit = [k for k in by_param if k.upper() in (dic or '').upper()]
    if not hit:
        hit = [k for k in by_param if k.upper() == name.upper().replace(' ', '')]
    if hit:
        k = hit[0]
        p('| %s | %s | `%s` | %s | `%s` | есть |' % (sect, name, dic, by_param[k][0], by_param[k][1]))
    else:
        p('| %s | %s | `%s` | — | — | **нет комбо** |' % (sect, name, dic))
p()
p('### Комбо формы, которых нет среди «Условий выборки» ТТ')
p()
for k in param_names:
    if not any(k.upper() in (d or '').upper() for _, _, d in filters):
        p('- `%s` (%s, атрибут `%s`)' % (k, by_param[k][0], by_param[k][1]))
p()

p('## 2. «Поля для вывода» (ТТ) ↔ флажки формы (P_FIELD_LIST)')
p()
cb_by_field = {}
for cb, fld in cb_fields:
    cb_by_field.setdefault(fld, []).append(cb)
p('| Раздел | Поле (ТТ) | По умолчанию (ТТ) | Флажок в форме | Статус |')
p('|---|---|---|---|---|')
for sect, name, dflt in outputs:
    key = name.upper().replace(' ', '').replace('.', '').replace('(', '').replace(')', '')
    hit = [f for f in cb_by_field if f.upper().replace('_', '') == key or key.startswith(f.upper().replace('_', '')) or f.upper().replace('_', '').startswith(key)]
    if hit:
        p('| %s | %s | %s | %s | есть |' % (sect, name, dflt or '', ', '.join(cb_by_field[hit[0]])))
    else:
        p('| %s | %s | %s | — | **нет флажка/обработчика** |' % (sect, name, dflt or ''))
p()
p('### Поля в P_FIELD_LIST (код), которых нет в списке «Поля для вывода» ТТ')
p()
tt_out_words = ' | '.join(o[1].upper() for o in outputs)
for fld, cbs in sorted(cb_by_field.items()):
    if fld.upper() not in tt_out_words:
        p('- `%s` ← %s' % (fld, ', '.join(cbs)))
p()

p('## 3. Идентификаторы полей отчёта (лист «Вывод полей в отчете», колонки «поле»/«Поле») и параметры формы')
p()
ids = set()
for rf in report_fields:
    for x in (rf[3], rf[6]):
        if re.fullmatch(r'[A-Za-z][A-Za-z0-9_]*', x or ''):
            ids.add(x)
missing_in_code = sorted(i for i in ids if i.upper() not in {n.upper() for n in field_names} and i.upper() not in {n.upper() for n in param_names})
p('Идентификаторов в ТТ: %d. Нет ни в P_FIELD_LIST, ни среди параметров формы (%d):' % (len(ids), len(missing_in_code)))
for i in missing_in_code:
    p('- `%s`' % i)
p()
code_only_fields = sorted(set(field_names) | set(param_names))
p('Идентификаторы формы, которых нет в колонках «поле»/«Поле» листа «Вывод полей в отчете»:')
for i in code_only_fields:
    if i.upper() not in {x.upper() for x in ids}:
        p('- `%s`' % i)
p()

p('## 4. Верхний уровень СЭ (лист «Селекционный экран»)')
p()
p('| Поле СЭ | Описание | Справочник |')
p('|---|---|---|')
for a, b, c in screen:
    p('| %s | %s | %s |' % (a, b, c))
p()
p('## 5. Объявленные контролы формы (по типам)')
p()
types = {}
for n, t in declared:
    types.setdefault(t, []).append(n)
for t in sorted(types):
    p('- **%s** (%d): %s' % (t, len(types[t]), ', '.join(types[t])))

rep = os.path.join(outdir, 'tt-vs-code.md')
open(rep, 'w', encoding='utf-8').write('\n'.join(out) + '\n')
print('отчёт: %s (%d строк)' % (rep, len(out)))
